"""Read spreadsheets back into Python (and into ``Document`` objects).

Supports ``.xlsx``/``.xlsm`` (openpyxl), ``.xls`` (``xlrd``, ``marmir[xls]``),
``.csv``/``.tsv`` and JSON (workbook specs or lists of records).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO, Any

from .composers import format_for_path, normalise_format
from .errors import MissingDependencyError, UnsupportedFormatError

__all__ = ["Sheet", "read", "read_document"]

Source = str | os.PathLike[str] | bytes | IO[Any]


@dataclass
class Sheet:
    """One sheet of data read from a file."""

    name: str
    columns: list[str]
    rows: list[list[Any]] = field(default_factory=list)

    def records(self) -> list[dict[str, Any]]:
        return [dict(zip(self.columns, row)) for row in self.rows]

    def __len__(self) -> int:
        return len(self.rows)


def read(
    source: Source,
    format: str | None = None,
    *,
    header: bool = True,
    sheet: str | int | None = None,
    max_rows: int | None = None,
    formulas: bool = True,
) -> list[Sheet]:
    """Read every sheet (or just ``sheet``) from ``source``.

    Args:
        source: a path, raw bytes or a binary file object.
        format: ``"xlsx"``, ``"xls"``, ``"csv"``, ``"tsv"`` or ``"json"``; guessed from the
            file extension (or content) when omitted.
        header: treat the first row as column titles.
        sheet: a sheet name or 0-based index to read only one sheet.
        max_rows: stop after this many data rows per sheet.
        formulas: return formulas as ``"=SUM(...)"`` text (``True``) or the cached
            values stored in the file (``False``; XLSX only).
    """
    fmt = _detect_format(source, format)
    data = _read_bytes(source)
    if fmt == "xlsx":
        sheets = _read_xlsx(data, header, max_rows, formulas, sheet)
    elif fmt == "xls":
        sheets = _read_xls(data, header, max_rows, sheet)
    elif fmt in ("csv", "tsv"):
        sheets = [_read_csv(data, header, max_rows, "\t" if fmt == "tsv" else None)]
    elif fmt == "json":
        sheets = _read_json(data, max_rows)
    else:
        raise UnsupportedFormatError(f"cannot read format {fmt!r}")
    if sheet is not None and fmt not in ("xlsx", "xls"):
        sheets = _select(sheets, sheet)
    return sheets


def read_document(source: Source, format: str | None = None, **kwargs: Any) -> Any:
    """Read a file into a ``Document`` (extra sheets become children)."""
    from .document import Document
    from .spec import from_spec

    fmt = _detect_format(source, format)
    if fmt == "json":
        payload = json.loads(_read_bytes(source))
        if isinstance(payload, dict) and ("sheets" in payload or "rows" in payload):
            return from_spec(payload, allow_images=False)
    sheets = read(source, fmt, **kwargs)
    root: Document | None = None
    for item in sheets:
        doc = Document(item.rows, order=item.columns, name=item.name)
        if root is None:
            root = doc
        else:
            root.add_child(doc)
    if root is None:
        raise UnsupportedFormatError("no sheets found")
    return root


def _detect_format(source: Source, format: str | None) -> str:
    if format:
        return normalise_format(format)
    if isinstance(source, (str, os.PathLike)):
        guessed = format_for_path(source)
        if guessed:
            return guessed
    name = getattr(source, "name", None)
    if isinstance(name, str) and format_for_path(name):
        return format_for_path(name) or ""
    head = _peek(source)
    if head.startswith(b"PK"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if head.lstrip()[:1] in (b"{", b"["):
        return "json"
    return "csv"


def _peek(source: Source) -> bytes:
    if isinstance(source, bytes):
        return source[:8]
    if isinstance(source, (str, os.PathLike)):
        with open(source, "rb") as handle:
            return handle.read(8)
    if hasattr(source, "peek"):
        return bytes(source.peek(8)[:8])
    if hasattr(source, "seek") and hasattr(source, "tell"):
        pos = source.tell()
        head = source.read(8)
        source.seek(pos)
        return head if isinstance(head, bytes) else head.encode()
    return b""


def _read_bytes(source: Source) -> bytes:
    if isinstance(source, bytes):
        return source
    if isinstance(source, (str, os.PathLike)):
        return Path(source).read_bytes()
    data = source.read()
    return data if isinstance(data, bytes) else str(data).encode("utf-8")


def _split(
    rows: list[list[Any]], header: bool, max_rows: int | None
) -> tuple[list[str], list[list[Any]]]:
    while rows and all(v is None or v == "" for v in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    rows = [list(r) + [None] * (width - len(r)) for r in rows]
    if header and rows:
        columns = [
            str(c) if c not in (None, "") else f"Column {i + 1}" for i, c in enumerate(rows[0])
        ]
        body = rows[1:]
    else:
        columns = [f"Column {i + 1}" for i in range(width)]
        body = rows
    if max_rows is not None:
        body = body[:max_rows]
    return columns, body


def _select(sheets: list[Sheet], sheet: str | int) -> list[Sheet]:
    if isinstance(sheet, int):
        return [sheets[sheet]]
    for item in sheets:
        if item.name == sheet:
            return [item]
    raise KeyError(f"no sheet named {sheet!r}; sheets are {[s.name for s in sheets]}")


def _read_xlsx(
    data: bytes, header: bool, max_rows: int | None, formulas: bool, sheet: str | int | None
) -> list[Sheet]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), data_only=not formulas, read_only=True)
    try:
        worksheets = workbook.worksheets
        if sheet is not None:
            worksheets = (
                [workbook.worksheets[sheet]] if isinstance(sheet, int) else [workbook[sheet]]
            )
        out = []
        for ws in worksheets:
            limit = None if max_rows is None else max_rows + (1 if header else 0)
            rows = []
            for row in ws.iter_rows():
                rows.append([_cell_value(c) for c in row])
                if limit is not None and len(rows) >= limit:
                    break
            columns, body = _split(rows, header, max_rows)
            out.append(Sheet(ws.title, columns, body))
        return out
    finally:
        workbook.close()


def _cell_value(cell: Any) -> Any:
    value = getattr(cell, "value", None)
    if isinstance(value, dt.datetime) and value.time() == dt.time(0):
        fmt = str(getattr(cell, "number_format", "") or "").lower()
        if not any(code in fmt for code in ("h", "s", "am/pm")):
            return value.date()  # a date-only format: hand back a date
    return value


def _read_xls(
    data: bytes, header: bool, max_rows: int | None, sheet: str | int | None
) -> list[Sheet]:
    try:
        import xlrd
    except ImportError as exc:
        raise MissingDependencyError("xlrd", "xls", "Reading .xls files") from exc
    book = xlrd.open_workbook(file_contents=data)
    indexes = range(book.nsheets)
    if sheet is not None:
        indexes = [sheet] if isinstance(sheet, int) else [book.sheet_names().index(sheet)]  # type: ignore[assignment]
    out = []
    for index in indexes:
        ws = book.sheet_by_index(index)
        rows = []
        for r in range(ws.nrows):
            row = []
            for c in range(ws.ncols):
                cell = ws.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                elif cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
                    row.append(int(cell.value))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        columns, body = _split(rows, header, max_rows)
        out.append(Sheet(ws.name, columns, body))
    return out


def _read_csv(data: bytes, header: bool, max_rows: int | None, delimiter: str | None) -> Sheet:
    text = data.decode("utf-8-sig")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    rows = [
        [_from_text(v) for v in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)
    ]
    columns, body = _split(rows, header, max_rows)
    return Sheet("Sheet 1", columns, body)


def _from_text(value: str) -> Any:
    if value == "":
        return None
    if value.startswith("'") and len(value) > 1 and value[1] in "=+-@":
        return value[1:]  # undo formula-injection escaping
    for cast in (int, float):
        try:
            return cast(value)
        except ValueError:
            pass
    return value


def _read_json(data: bytes, max_rows: int | None) -> list[Sheet]:
    from .spec import from_spec

    payload = json.loads(data)
    document = from_spec(payload, allow_images=False)
    out = []
    for doc in document.sheets:
        rows = [[_plain(cell) for cell in row] for row in doc.grid.grid_data]
        if max_rows is not None:
            rows = rows[:max_rows]
        out.append(
            Sheet(doc.name or "Sheet 1", [str(t) for t in doc.data_model.field_titles], rows)
        )
    return out


def _plain(cell: Any) -> Any:
    from .types import FormulaFieldType

    if isinstance(cell, FormulaFieldType):
        return "=" + cell.expression
    return cell.data
