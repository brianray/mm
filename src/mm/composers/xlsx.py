"""Excel 2007+ (``.xlsx``) output via openpyxl. This is the default format."""

from __future__ import annotations

import copy
import io
import logging
import re
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..errors import MissingDependencyError
from ..fonts import string_width_chars
from ..style import Style
from ..types import (
    BaseFieldType,
    DateTimeFieldType,
    FormulaFieldType,
    ImageFieldType,
    NoneFieldType,
    StringFieldType,
    URLFieldType,
)
from .base import ComposerBase, cell_style, column_key_lookup, header_style, number_format

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

    from ..document import Document

log = logging.getLogger(__name__)

__all__ = ["ComposerXLSX", "safe_sheet_name"]

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_LINK_COLOR = "0563C1"


def safe_sheet_name(name: str | None, index: int, used: set[str]) -> str:
    """Excel sheet names: max 31 chars, no ``[]:*?/\\``, unique (case-insensitive)."""
    base = (
        _INVALID_SHEET_CHARS.sub("_", str(name or f"Sheet {index}")).strip("'") or f"Sheet {index}"
    )
    base = base[:31]
    candidate = base
    counter = 2
    while candidate.lower() in used:
        suffix = f" ({counter})"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(candidate.lower())
    return candidate


class ComposerXLSX(ComposerBase):
    format_name = "xlsx"
    binary = True
    media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def __init__(self, data_model: Any, grid: Any, document: Document):
        super().__init__(data_model, grid, document)
        self._style_cache: dict[tuple[Style, str | None, bool], Any] = {}

    def run(self, child: Any = None) -> bytes | None:
        top = child is None
        workbook: Workbook = Workbook() if top else child
        if top:
            workbook.remove(workbook.active)
            self._set_properties(workbook)
        used = {ws.title.lower() for ws in workbook.worksheets}
        for document in self.sheets():
            title = safe_sheet_name(document.name, len(workbook.worksheets) + 1, used)
            worksheet = workbook.create_sheet(title)
            self.write_sheet(worksheet, document)
        if not top:
            return None
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def _set_properties(self, workbook: Workbook) -> None:
        config = self.document.config
        props = workbook.properties
        props.title = config.get("title") or self.document.name
        props.creator = config.get("author") or "Marmir"
        props.description = config.get("description")

    # -- one sheet -------------------------------------------------------------------
    def write_sheet(self, ws: Worksheet, document: Document) -> None:
        config = document.config
        titles = document.data_model.field_titles
        widths: dict[int, float] = {}
        longest: dict[int, int] = {}
        row_offset = 0

        if config.get("headers", True) and titles:
            style = header_style(document)
            for col, title in enumerate(titles):
                cell = ws.cell(row=1, column=col + 1, value=_clean(str(title)))
                self._apply_style(cell, style, None)
                widths[col] = self._measure(str(title), style)
            row_offset = 1

        for data_row, row in enumerate(document.grid.grid_data):
            for col, field in enumerate(row):
                target = ws.cell(row=data_row + row_offset + 1, column=col + 1)
                style = cell_style(document, data_row, col, field)
                fmt = number_format(document, col, field, style)
                text = self._write_value(ws, target, field)
                self._apply_style(target, style, fmt, link=isinstance(field, URLFieldType))
                # only strings near the longest seen so far can widen the column
                if text and len(text) * 1.5 >= longest.get(col, 0):
                    longest[col] = max(longest.get(col, 0), len(text))
                    widths[col] = max(widths.get(col, 0.0), self._measure(text, style))
                if isinstance(field, ImageFieldType):
                    self._place_image(ws, target, field, col, widths)

        self._finish_sheet(ws, document, widths, row_offset)

    def _write_value(self, ws: Worksheet, target: Any, field: BaseFieldType) -> str:
        """Set the cell value; return the text used to size the column."""
        if isinstance(field, NoneFieldType) or field.data is None:
            return ""
        if isinstance(field, FormulaFieldType):
            target.value = "=" + field.expression
            return ""
        if isinstance(field, URLFieldType):
            target.value = _clean(str(field.displayname))
            target.hyperlink = str(field.data)
            return str(field.displayname)
        if isinstance(field, ImageFieldType):
            return ""
        if isinstance(field, StringFieldType) or isinstance(field.data, str):
            text = _clean(str(field.data))
            target.value = text
            if text.startswith("="):
                target.data_type = "s"  # plain text, never a formula
            return text
        target.value = field.data
        if isinstance(field, DateTimeFieldType):
            return "0000-00-00 00:00:00"
        return field.display()

    def _place_image(
        self, ws: Worksheet, target: Any, field: ImageFieldType, col: int, widths: dict[int, float]
    ) -> None:
        try:
            from openpyxl.drawing.image import Image as XLImage
        except ImportError as exc:  # pragma: no cover - openpyxl always ships this
            raise MissingDependencyError("pillow", "images", "Images in XLSX") from exc
        try:
            image = XLImage(str(field.data))
        except ImportError as exc:
            raise MissingDependencyError("pillow", "images", "Images in XLSX") from exc
        if field.width:
            image.width = field.width
        if field.height:
            image.height = field.height
        ws.add_image(image, target.coordinate)
        widths[col] = max(widths.get(col, 0.0), image.width / 7.0)
        current = ws.row_dimensions[target.row].height or 15
        ws.row_dimensions[target.row].height = max(current, image.height * 0.75)

    def _finish_sheet(
        self, ws: Worksheet, document: Document, widths: dict[int, float], row_offset: int
    ) -> None:
        config = document.config
        titles = document.data_model.field_titles
        freeze_row = int(config.get("freeze_row", 0) or 0)
        freeze_col = int(config.get("freeze_col", 0) or 0)
        if not config.get("headers", True) and freeze_row == 1:
            freeze_row = 0  # nothing to freeze
        if freeze_row > 0 or freeze_col > 0:
            ws.freeze_panes = f"{get_column_letter(freeze_col + 1)}{freeze_row + 1}"

        if config.get("autofilter") and row_offset and titles:
            last_row = max(1, len(document.grid.grid_data) + row_offset)
            ws.auto_filter.ref = f"A1:{get_column_letter(len(titles))}{last_row}"

        fixed = config.get("column_widths") or {}
        auto = config.get("adjust_all_col_width", True)
        lo = float(config.get("min_col_width", 0) or 0)
        hi = float(config.get("max_col_width", 255) or 255)
        for col, title in enumerate(titles):
            letter = get_column_letter(col + 1)
            width = column_key_lookup(fixed, title, col)
            if width is None and auto and col in widths:
                width = min(hi, max(lo, widths[col] + 2))
            if width is not None:
                ws.column_dimensions[letter].width = float(width)

    # -- styles ----------------------------------------------------------------------
    def _measure(self, text: str, style: Style) -> float:
        return string_width_chars(text, style.font_family, style.font_size or 11, bool(style.bold))

    def _apply_style(self, cell: Any, style: Style, fmt: str | None, link: bool = False) -> None:
        key = (style, fmt, link)
        cached = self._style_cache.get(key)
        if cached is None:
            # first cell with this style: let openpyxl register the style objects, then
            # remember the resulting style array -- assigning font/fill/... on every cell
            # makes openpyxl re-hash them and dominates write time on large sheets
            for attr, value in self._build_style(style, fmt, link).items():
                setattr(cell, attr, value)
            self._style_cache[key] = copy.copy(cell._style)
        else:
            cell._style = copy.copy(cached)

    @staticmethod
    def _build_style(style: Style, fmt: str | None, link: bool) -> dict[str, Any]:
        out: dict[str, Any] = {}
        color = style.color
        underline = "single" if style.underline else None
        if link:
            color = color if color and color != "000000" else _LINK_COLOR
            underline = "single"
        if any(
            v is not None
            for v in (
                style.font_family,
                style.font_size,
                style.bold,
                style.italic,
                style.strike,
                color,
                underline,
            )
        ):
            out["font"] = Font(
                name=style.font_family or "Arial",
                size=style.font_size or 10,
                bold=bool(style.bold),
                italic=bool(style.italic),
                strike=bool(style.strike),
                underline=underline,
                color=color,
            )
        if style.background_color:
            out["fill"] = PatternFill(
                fill_type="solid",
                start_color=style.background_color,
                end_color=style.background_color,
            )
        if style.text_align or style.vertical_align or style.wrap is not None:
            out["alignment"] = Alignment(
                horizontal=style.text_align,
                vertical=style.vertical_align,
                wrap_text=style.wrap,
            )
        if style.border_style and style.border_style != "none":
            side = Side(style=style.border_style, color=style.border_color or "000000")
            out["border"] = Border(left=side, right=side, top=side, bottom=side)
        elif style.border_color:
            side = Side(style="thin", color=style.border_color)
            out["border"] = Border(left=side, right=side, top=side, bottom=side)
        if fmt:
            out["number_format"] = fmt
        return out


def _clean(text: str) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", text)
