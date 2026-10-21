"""Helpers to read generated workbooks back."""

from __future__ import annotations

import io

from openpyxl import load_workbook


def load_xlsx(data: bytes):
    return load_workbook(io.BytesIO(data))


def sheet_values(data: bytes, index: int = 0) -> list[list]:
    ws = load_xlsx(data).worksheets[index]
    return [list(row) for row in ws.iter_rows(values_only=True)]
