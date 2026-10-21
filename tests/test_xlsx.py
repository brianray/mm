from __future__ import annotations

import datetime as dt
from decimal import Decimal

import pytest

import mm

from .helpers import load_xlsx


def _ws(doc, index=0):
    return load_xlsx(doc.writestr("xlsx")).worksheets[index]


def test_types_and_number_formats(now):
    doc = mm.Document(
        [
            {
                "int": 1,
                "float": 2.5,
                "dec": Decimal("3.14"),
                "bool": True,
                "date": dt.date(2024, 1, 2),
                "dt": now,
                "time": dt.time(13, 30),
                "dur": dt.timedelta(hours=2),
                "pct": mm.Percent(0.125, decimals=1),
                "money": mm.Currency(1234.5),
                "custom": mm.Date(now, "%Y-%m-%dT%H:%M:%S"),
                "none": None,
            }
        ]
    )
    ws = _ws(doc)
    row = {ws.cell(1, c).value: ws.cell(2, c) for c in range(1, ws.max_column + 1)}
    assert row["int"].value == 1
    assert row["float"].value == 2.5
    assert row["dec"].value == pytest.approx(3.14)
    assert row["bool"].value is True
    assert row["date"].value == dt.datetime(2024, 1, 2)
    assert row["date"].number_format == "M/D/YY"
    assert row["dt"].value == now
    assert row["dt"].number_format == "M/D/YY h:mm:ss"
    assert row["time"].value == dt.time(13, 30)
    assert row["dur"].number_format == "[h]:mm:ss"
    assert row["pct"].number_format == "0.0%"
    assert row["money"].number_format.startswith('"$"#,##0.00')
    assert row["custom"].number_format == 'yyyy-mm-dd"T"hh:mm:ss'
    assert row["none"].value is None


def test_column_number_formats_and_widths():
    doc = mm.Document(
        [{"amount": 1.5, "share": mm.Percent(0.5)}],
        number_formats={"amount": "#,##0.000", 1: "0.00%"},
        column_widths={"A": 33},
    )
    ws = _ws(doc)
    assert ws["A2"].number_format == "#,##0.000"
    assert ws["B2"].number_format == "0.00%"
    assert ws.column_dimensions["A"].width == 33


def test_formulas_urls_and_text_that_looks_like_formula():
    doc = mm.Document(
        [
            {"n": 1, "link": mm.URL("https://github.com/brianray", "GitHub"), "s": "x"},
            {"n": 2, "link": mm.URL("http://twitter.com/brianray"), "s": "=not a formula"},
            {"n": mm.Formula("SUM(A2:A3)"), "link": None, "s": "=A2*2"},
        ],
        column_types={"s": "string"},
    )
    ws = _ws(doc)
    assert ws["A4"].value == "=SUM(A2:A3)"
    assert ws["B2"].value == "GitHub"
    assert ws["B2"].hyperlink.target == "https://github.com/brianray"
    assert ws["B2"].font.underline == "single"
    assert ws["B3"].value == "http://twitter.com/brianray"
    assert ws["C3"].value == "=not a formula"
    assert ws["C3"].data_type == "s"


def test_inferred_formula_strings():
    ws = _ws(mm.Document([{"a": 1}, {"a": 2}, {"a": "=SUM(A2:A3)"}]))
    assert ws["A4"].value == "=SUM(A2:A3)"
    assert ws["A4"].data_type == "f"
    ws = _ws(mm.Document([{"a": "=SUM(A2:A3)"}], infer_formulas=False))
    assert ws["A2"].data_type == "s"


def test_styles_and_layout():
    doc = mm.Document(
        [
            {"a": 1, "b": mm.Cell("hi", style="background-color: yellow; font-style: italic")},
            {"a": 2, "b": "x"},
        ],
        autofilter=True,
        freeze_col=1,
        column_styles={"a": "text-align: right; border: 1px solid #999999"},
        title="T",
        author="Me",
    )
    wb = load_xlsx(doc.writestr())
    ws = wb.active
    header = ws["A1"]
    assert header.font.bold is True
    assert header.font.color.rgb.endswith("FFFFFF")
    assert header.fill.start_color.rgb.endswith("1F4E79")
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:B3"
    assert ws["B2"].fill.start_color.rgb.endswith("FFFF00")
    assert ws["B2"].font.italic is True
    assert ws["A2"].alignment.horizontal == "right"
    assert ws["A2"].border.left.style == "thin"
    assert ws["A3"].fill.start_color.rgb.endswith("F2F2F2")  # banding
    assert wb.properties.title == "T"
    assert wb.properties.creator == "Me"


def test_no_headers_no_styles():
    doc = mm.Document([{"a": 1}], headers=False, header_style="", row_styles=())
    ws = _ws(doc)
    assert ws["A1"].value == 1
    assert ws.freeze_panes is None
    assert ws["A1"].fill.fill_type is None


def test_auto_width_grows_with_content():
    doc = mm.Document([{"short": "a", "long": "a much much longer piece of text"}])
    ws = _ws(doc)
    assert ws.column_dimensions["B"].width > ws.column_dimensions["A"].width
    assert ws.column_dimensions["B"].width <= 60
    off = _ws(mm.Document([{"x": "long text " * 5}], adjust_all_col_width=False))
    assert off.column_dimensions["A"].width == 13  # openpyxl default


def test_multiple_sheets_names_sanitized():
    doc = mm.Document([{"a": 1}], name="Bad/Name:*?")
    doc.add_sheet([{"b": 2}], name="Bad/Name:*?")
    doc.add_sheet([{"c": 3}], name="x" * 40)
    doc.add_sheet([{"d": 4}])
    wb = load_xlsx(doc.writestr())
    assert wb.sheetnames == ["Bad_Name___", "Bad_Name___ (2)", "x" * 31, "Sheet 4"]
    assert wb.worksheets[1]["A2"].value == 2


def test_images(image_path):
    pytest.importorskip("PIL")
    doc = mm.Document([{"profile": mm.Image(str(image_path), 230, 326)}])
    ws = _ws(doc)
    assert len(ws._images) == 1
    assert ws.row_dimensions[2].height == pytest.approx(326 * 0.75)


def test_illegal_characters_are_stripped():
    ws = _ws(mm.Document([{"a": "bad\x01char"}]))
    assert ws["A2"].value == "badchar"


def test_composer_into_existing_workbook(records):
    from openpyxl import Workbook

    from mm.composers.xlsx import ComposerXLSX

    wb = Workbook()
    doc = mm.Document(records, name="Extra")
    assert ComposerXLSX(doc.data_model, doc.grid, doc).run(child=wb) is None
    assert "Extra" in wb.sheetnames


def test_column_number_format_applies_to_formulas():
    doc = mm.Document(
        [{"qty": 2, "total": mm.Formula("A2*3")}, {"qty": 1, "total": "=A3*3"}],
        number_formats={"total": '"$"#,##0.00'},
    )
    ws = _ws(doc)
    assert ws["B2"].number_format == '"$"#,##0.00'
    assert ws["B3"].number_format == '"$"#,##0.00'
